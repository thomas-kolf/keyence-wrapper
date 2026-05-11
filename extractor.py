from pathlib import Path
from dataclasses import dataclass, asdict
from openpyxl import load_workbook
import json

from config_loader import machine_config


EXCEL_METADATA_CONFIG = machine_config["excel_metadata"]
MEASUREMENT_TABLE_CONFIG = machine_config["measurement_table"]
MEASUREMENT_COLUMNS_CONFIG = machine_config["measurement_table"]["columns"]
QUALITY_MAPPING_CONFIG = machine_config["quality_mapping"]


@dataclass
class CellMetadata:
    source_file: str
    timestamp: str | None
    leadframe_dmc: str | None
    position: str | None
    name: str | None
    product_name: str | None
    device: str | None
    overall_result: str | None


@dataclass
class Measurement:
    nr: str | None
    measurement_name: str | None
    elem_1: str | None
    detail: str | None
    elem_2: str | None
    kommentar: str | None
    classification: str | None
    value: str | None
    unit: str | None
    target: str | None
    upper_tolerance: str | None
    lower_tolerance: str | None


def find_excel_files(file_group: list[Path]) -> list[Path]:
    return [
        file_path
        for file_path in file_group
        if file_path.suffix.lower() == ".xlsx"
    ]


def clean_cell_value(value) -> str | None:
    if value is None:
        return None

    value = str(value).strip()

    if value == "":
        return None

    return value


def read_metadata_from_excel(excel_file: Path) -> CellMetadata:
    workbook = load_workbook(excel_file, data_only=True, read_only=True)

    try:
        sheet = workbook.active

        return CellMetadata(
            source_file=excel_file.name,
            timestamp=clean_cell_value(
                sheet[EXCEL_METADATA_CONFIG["timestamp_cell"]].value
            ),
            leadframe_dmc=clean_cell_value(
                sheet[EXCEL_METADATA_CONFIG["dmc_cell"]].value
            ),
            position=clean_cell_value(
                sheet[EXCEL_METADATA_CONFIG["position_cell"]].value
            ),
            name=clean_cell_value(
                sheet[EXCEL_METADATA_CONFIG["name_cell"]].value
            ),
            product_name=clean_cell_value(
                sheet[EXCEL_METADATA_CONFIG["product_name_cell"]].value
            ),
            device=clean_cell_value(
                sheet[EXCEL_METADATA_CONFIG["device_cell"]].value
            ),
            overall_result=clean_cell_value(
                sheet[EXCEL_METADATA_CONFIG["overall_result_cell"]].value
            ),
        )

    finally:
        workbook.close()


def extract_metadata(file_group: list[Path]) -> list[CellMetadata]:
    excel_files = find_excel_files(file_group)

    metadata_list = []

    for excel_file in excel_files:
        metadata = read_metadata_from_excel(excel_file)
        metadata_list.append(metadata)

    return metadata_list


def normalize_positions(metadata_list: list[CellMetadata]) -> dict[str, str]:
    raw_positions = []

    for metadata in metadata_list:
        if metadata.position is not None:
            raw_positions.append(metadata.position)

    sorted_positions = sorted(raw_positions, key=lambda pos: int(pos))

    position_mapping = {}

    for index, raw_position in enumerate(sorted_positions, start=1):
        normalized_position = f"{index:02d}"
        position_mapping[raw_position] = normalized_position

    return position_mapping


def read_measurement_cell(sheet, field_name: str, row: int) -> str | None:
    column = MEASUREMENT_COLUMNS_CONFIG[field_name]
    return clean_cell_value(sheet[f"{column}{row}"].value)


def extract_measurements_from_excel(excel_file: Path) -> list[Measurement]:
    workbook = load_workbook(excel_file, data_only=True, read_only=True)

    measurements = []

    try:
        sheet = workbook.active
        row = MEASUREMENT_TABLE_CONFIG["start_row"]

        while True:
            nr = read_measurement_cell(sheet, "nr", row)

            if nr is None:
                break

            measurement = Measurement(
                nr=nr,
                measurement_name=read_measurement_cell(
                    sheet, "measurement_name", row
                ),
                elem_1=read_measurement_cell(sheet, "elem_1", row),
                detail=read_measurement_cell(sheet, "detail", row),
                elem_2=read_measurement_cell(sheet, "elem_2", row),
                kommentar=read_measurement_cell(sheet, "kommentar", row),
                classification=read_measurement_cell(sheet, "classification", row),
                value=read_measurement_cell(sheet, "value", row),
                unit=read_measurement_cell(sheet, "unit", row),
                target=read_measurement_cell(sheet, "target", row),
                upper_tolerance=read_measurement_cell(
                    sheet, "upper_tolerance", row
                ),
                lower_tolerance=read_measurement_cell(
                    sheet, "lower_tolerance", row
                ),
            )

            measurements.append(measurement)
            row += 1

    finally:
        workbook.close()

    return measurements


def determine_quality(metadata: CellMetadata) -> str:
    ok_value = QUALITY_MAPPING_CONFIG["ok_value"]
    ok_quality = QUALITY_MAPPING_CONFIG["ok_quality"]
    not_ok_quality = QUALITY_MAPPING_CONFIG["not_ok_quality"]

    if (
        metadata.overall_result is not None
        and metadata.overall_result.lower() == ok_value.lower()
    ):
        return ok_quality

    return not_ok_quality


def build_cell_data(file_group: list[Path]) -> list[dict]:
    metadata_list = extract_metadata(file_group)
    position_mapping = normalize_positions(metadata_list)

    cell_data_list = []

    for metadata in metadata_list:
        normalized_position = position_mapping[metadata.position]
        cell_dmc = f"{metadata.leadframe_dmc}-{normalized_position}"

        excel_file = next(
            file for file in file_group
            if file.name == metadata.source_file
        )

        measurements = extract_measurements_from_excel(excel_file)
        quality = determine_quality(metadata)

        cell_data = {
            "source_file": metadata.source_file,
            "timestamp": metadata.timestamp,
            "leadframe_dmc": metadata.leadframe_dmc,
            "raw_position": metadata.position,
            "position": normalized_position,
            "cell_dmc": cell_dmc,
            "name": metadata.name,
            "product_name": metadata.product_name,
            "device": metadata.device,
            "overall_result": metadata.overall_result,
            "quality": quality,
            "measurements": [
                asdict(measurement)
                for measurement in measurements
            ],
        }

        cell_data_list.append(cell_data)

    return cell_data_list


def write_cell_json(cell_data: dict, output_dir: Path, file_base_name: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    output_file = output_dir / f"{file_base_name}.json"

    with output_file.open("w", encoding="utf-8") as json_file:
        json.dump(cell_data, json_file, ensure_ascii=False, indent=4)

    return output_file